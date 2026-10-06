"""Text extraction from uploaded resumes (PDF, DOCX, TXT)."""

import io

from docx import Document
from pypdf import PdfReader
from pypdf.errors import PdfReadError


class ParseError(ValueError):
    pass


def extract_text(data: bytes, ext: str) -> str:
    ext = ext.lower()
    if ext == ".txt":
        for encoding in ("utf-8", "utf-16", "latin-1"):
            try:
                return data.decode(encoding)
            except UnicodeDecodeError:
                continue
        raise ParseError("Could not decode text file.")
    if ext == ".pdf":
        try:
            reader = PdfReader(io.BytesIO(data))
            return "\n".join(page.extract_text() or "" for page in reader.pages)
        except (PdfReadError, ValueError, KeyError) as exc:
            raise ParseError(f"Could not read PDF: {exc}") from exc
    if ext == ".docx":
        try:
            doc = Document(io.BytesIO(data))
        except Exception as exc:  # python-docx raises several unrelated types on bad input
            raise ParseError(f"Could not read DOCX: {exc}") from exc
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        return "\n".join(parts)
    raise ParseError(f"Unsupported file type: {ext}")
